"""Turn a file into something a model can read: text, images, and hints."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageOps

try:
    import pillow_heif

    pillow_heif.register_heif_opener()
except ImportError:  # HEIC support is optional
    pass

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".heic", ".heif"}
MIN_PDF_TEXT = 40  # fewer characters than this means "scanned PDF"


class ExtractionError(Exception):
    pass


@dataclass
class Content:
    text: str = ""
    images: list[bytes] = field(default_factory=list)  # JPEG bytes
    hints: dict = field(default_factory=dict)

    @property
    def empty(self) -> bool:
        return not self.text.strip() and not self.images


def extract(path: Path, max_chars: int = 6000, image_max_side: int = 1024, pdf_scan_pages: int = 2) -> Content:
    ext = path.suffix.lower()
    content = Content(hints={
        "current_filename": path.name,
        "file_modified": datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d"),
    })
    try:
        if ext == ".pdf":
            _pdf(path, content, max_chars, image_max_side, pdf_scan_pages)
        elif ext in IMAGE_EXT:
            _image(path, content, image_max_side)
        elif ext == ".docx":
            content.text = _docx(path)
        elif ext == ".xlsx":
            content.text = _xlsx(path, max_chars)
        elif ext == ".pptx":
            content.text = _pptx(path)
        else:
            content.text = _plain(path, max_chars)
    except ExtractionError:
        raise
    except Exception as e:  # corrupt or password-protected files, etc.
        raise ExtractionError(f"Could not read {ext} file: {e}") from e
    content.text = _clean(content.text)[:max_chars]
    return content


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    text = "".join(ch for ch in text if ch.isprintable() or ch == "\n")
    text = re.sub(r" {2,}", " ", text)
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


def _to_jpeg(img: Image.Image, max_side: int) -> bytes:
    img = ImageOps.exif_transpose(img)
    if img.mode not in ("RGB", "L"):
        background = Image.new("RGB", img.size, "white")
        rgba = img.convert("RGBA")
        background.paste(rgba, mask=rgba.split()[-1])
        img = background
    img.thumbnail((max_side, max_side))
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "JPEG", quality=85)
    return buf.getvalue()


def _image(path: Path, content: Content, max_side: int) -> None:
    with Image.open(path) as img:
        exif = img.getexif()
        # 36867 = DateTimeOriginal (in the Exif IFD), 306 = DateTime
        taken = exif.get_ifd(0x8769).get(36867) or exif.get(306)
        if taken:
            content.hints["photo_taken"] = str(taken)[:10].replace(":", "-")
        content.hints["image_size"] = f"{img.width}x{img.height}"
        content.images.append(_to_jpeg(img, max_side))


def _pdf(path: Path, content: Content, max_chars: int, max_side: int, scan_pages: int) -> None:
    import pymupdf

    with pymupdf.open(path) as doc:
        if doc.needs_pass:
            raise ExtractionError("PDF is password-protected")
        content.hints["pages"] = doc.page_count
        meta_title = (doc.metadata or {}).get("title")
        if meta_title:
            content.hints["pdf_title"] = meta_title
        parts, total = [], 0
        for page in doc:
            t = page.get_text()
            parts.append(t)
            total += len(t)
            if total >= max_chars:
                break
        text = "\n".join(parts)
        content.text = text
        if len(text.strip()) < MIN_PDF_TEXT:
            # Scanned PDF: let a vision model read the page images instead of OCR.
            for page in list(doc)[:scan_pages]:
                pix = page.get_pixmap(dpi=150)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                content.images.append(_to_jpeg(img, max(max_side, 1400)))


def _docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables[:5]:
        for row in table.rows[:30]:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    props = doc.core_properties
    if props.title:
        lines.insert(0, f"[Document title property: {props.title}]")
    return "\n".join(lines)


def _xlsx(path: Path, max_chars: int) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines, total = [], 0
    try:
        for ws in wb.worksheets[:3]:
            lines.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(max_row=60, values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    line = " | ".join(cells)
                    lines.append(line)
                    total += len(line)
                if total > max_chars:
                    return "\n".join(lines)
    finally:
        wb.close()
    return "\n".join(lines)


def _pptx(path: Path) -> str:
    from pptx import Presentation

    lines = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        lines.append(f"[Slide {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


def _plain(path: Path, max_chars: int) -> str:
    raw = path.read_bytes()[: max_chars * 4]
    if b"\x00" in raw[:2000] and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise ExtractionError("Looks like a binary file")
    encodings = ["utf-16"] if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "cp1252"]
    text = None
    for enc in encodings:
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError as e:
            if enc == "utf-8-sig" and e.start >= len(raw) - 4:  # multibyte char cut by the read limit
                text = raw[: e.start].decode(enc)
                break
    if text is None:
        text = raw.decode("latin-1")
    if path.suffix.lower() in (".html", ".htm", ".xml"):
        text = re.sub(r"(?is)<(script|style).*?</\1>", " ", text)
        text = re.sub(r"<[^>]+>", " ", text)
    elif path.suffix.lower() == ".rtf":
        text = re.sub(r"\\[a-z]+-?\d* ?|[{}]", " ", text)
    return text
