import tempfile
import unittest
from pathlib import Path

from PIL import Image

from ai_renamer.content import ExtractionError, extract


class ExtractTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_txt_utf8_and_cp1252(self):
        (self.dir / "a.txt").write_text("Recette: tajine épicé", encoding="utf-8")
        (self.dir / "b.txt").write_bytes("Café crème".encode("cp1252"))
        self.assertIn("épicé", extract(self.dir / "a.txt").text)
        self.assertIn("Café", extract(self.dir / "b.txt").text)

    def test_utf8_cut_mid_character_by_limit(self):
        (self.dir / "c.txt").write_text("é" * 100, encoding="utf-8")
        self.assertTrue(extract(self.dir / "c.txt", max_chars=25).text.startswith("ééé"))

    def test_html_tags_stripped(self):
        (self.dir / "p.html").write_text("<html><style>x{}</style><h1>Hello</h1><script>bad()</script></html>")
        text = extract(self.dir / "p.html").text
        self.assertIn("Hello", text)
        self.assertNotIn("bad", text)

    def test_binary_rejected(self):
        (self.dir / "x.log").write_bytes(b"\x00\x01\x02" * 100)
        with self.assertRaises(ExtractionError):
            extract(self.dir / "x.log")

    def test_docx_xlsx_pptx(self):
        import docx
        import openpyxl
        from pptx import Presentation

        d = docx.Document()
        d.add_paragraph("Rental agreement between A and B")
        d.save(self.dir / "a.docx")
        wb = openpyxl.Workbook()
        wb.active.append(["Budget", 2025])
        wb.save(self.dir / "b.xlsx")
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Quarterly review"
        prs.save(self.dir / "c.pptx")
        self.assertIn("Rental agreement", extract(self.dir / "a.docx").text)
        self.assertIn("Budget | 2025", extract(self.dir / "b.xlsx").text)
        self.assertIn("Quarterly review", extract(self.dir / "c.pptx").text)

    def test_pdf_text_and_scanned(self):
        import pymupdf

        doc = pymupdf.open()
        doc.new_page().insert_text((72, 72), "Invoice number INV-2041 from ACME Corporation, total due 1250 EUR")
        doc.save(self.dir / "text.pdf")
        doc.close()
        scan = pymupdf.open()
        page = scan.new_page()
        img = self.dir / "scan.png"
        Image.new("RGB", (200, 100), "white").save(img)
        page.insert_image(page.rect, filename=str(img))
        scan.save(self.dir / "scan.pdf")
        scan.close()

        text = extract(self.dir / "text.pdf")
        self.assertIn("INV-2041", text.text)
        self.assertEqual(len(text.images), 0)
        scanned = extract(self.dir / "scan.pdf")
        self.assertEqual(len(scanned.images), 1)
        self.assertEqual(scanned.hints["pages"], 1)

    def test_image_is_shrunk_and_exif_date_read(self):
        img = Image.new("RGBA", (3000, 1500), (255, 0, 0, 128))
        exif = Image.Exif()
        exif[306] = "2024:07:12 18:30:00"
        img.convert("RGB").save(self.dir / "p.jpg", exif=exif)
        img.save(self.dir / "q.png")
        c = extract(self.dir / "p.jpg", image_max_side=512)
        self.assertEqual(c.hints["photo_taken"], "2024-07-12")
        with Image.open(__import__("io").BytesIO(c.images[0])) as out:
            self.assertEqual(max(out.size), 512)
        self.assertEqual(len(extract(self.dir / "q.png").images), 1)  # transparency handled


if __name__ == "__main__":
    unittest.main()
